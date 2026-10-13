"""Structured local parsers preserve format-native positions."""

from __future__ import annotations

from pathlib import Path

import pytest

from searchos.local.models import (
    CellRangeLocator,
    DocxBlockLocator,
    DomRangeLocator,
    JsonPathLocator,
    PipelineOptions,
    PptxMode,
    RecordRangeLocator,
    SlideObjectLocator,
)
from searchos.local.errors import LocalDocumentError, UnsupportedDocumentError
from searchos.local.parsers import (
    DelimitedParser,
    DocxParser,
    HtmlParser,
    JsonLinesParser,
    JsonParser,
    ParserRegistry,
    PptxParser,
    XlsxParser,
)


class FakeSlideRenderer:
    def __init__(self) -> None:
        self.slide_indices: tuple[int, ...] = ()

    def render_slides(self, path, *, slide_indices, dpi, max_dimension):
        self.slide_indices = slide_indices
        return {index: "data:image/jpeg;base64,dGVzdA==" for index in slide_indices}


class FakeSlideVision:
    def __init__(self) -> None:
        self.native_text: dict[int, str] = {}

    async def transcribe_slide(self, *, slide_number, image_data_uri, native_text) -> str:
        assert image_data_uri.startswith("data:image/jpeg;base64,")
        self.native_text[slide_number] = native_text
        return f"第 {slide_number} 页视觉正文"


@pytest.mark.asyncio
async def test_csv_uses_logical_records_when_fields_contain_newlines(tmp_path: Path) -> None:
    source = tmp_path / "records.csv"
    source.write_text('name,note\nA,"first\nsecond"\nB,\n', encoding="utf-8")

    result = await DelimitedParser().parse(
        source,
        options=PipelineOptions(tabular_chunk_rows=1),
    )

    assert result.total_units == result.covered_units == 2
    assert len(result.segments) == 2
    first = result.segments[0].source_locators[0]
    second = result.segments[1].source_locators[0]
    assert isinstance(first, RecordRangeLocator)
    assert (first.record_start, first.record_end) == (1, 1)
    assert isinstance(second, RecordRangeLocator)
    assert (second.record_start, second.record_end) == (2, 2)
    assert "first\nsecond" in result.segments[0].text


@pytest.mark.asyncio
async def test_json_uses_escaped_rfc6901_paths(tmp_path: Path) -> None:
    source = tmp_path / "records.json"
    source.write_text('{"a/b": {"value": 1}, "a~b": null}', encoding="utf-8")

    result = await JsonParser().parse(source, options=PipelineOptions())

    paths = [segment.source_locators[0] for segment in result.segments]
    assert paths == [JsonPathLocator(path="/a~1b"), JsonPathLocator(path="/a~0b")]


@pytest.mark.asyncio
async def test_jsonl_uses_logical_record_locations(tmp_path: Path) -> None:
    source = tmp_path / "tasks.jsonl"
    source.write_text('{"id": 1}\n\n{"id": 2}\n', encoding="utf-8")

    result = await JsonLinesParser().parse(source, options=PipelineOptions())

    assert result.total_units == 2
    assert [segment.source_locators[0] for segment in result.segments] == [
        RecordRangeLocator(record_start=1, record_end=1),
        RecordRangeLocator(record_start=2, record_end=2),
    ]


@pytest.mark.asyncio
async def test_html_does_not_execute_or_include_scripts(tmp_path: Path) -> None:
    source = tmp_path / "page.html"
    source.write_text(
        "<html><body><script>fetch('https://example.com')</script>"
        "<h1>标题</h1><p>正文 <a href='https://example.com'>链接文字</a></p>"
        "</body></html>",
        encoding="utf-8",
    )

    result = await HtmlParser().parse(source, options=PipelineOptions())

    assert [segment.text for segment in result.segments] == ["标题", "正文 链接文字"]
    assert all(
        isinstance(segment.source_locators[0], DomRangeLocator) for segment in result.segments
    )
    assert result.parser_metadata["external_resources_loaded"] is False
    assert result.warnings


def test_registry_returns_explicit_unsupported_format(tmp_path: Path) -> None:
    source = tmp_path / "data.sqlite"
    source.write_bytes(b"SQLite format 3\x00")
    registry = ParserRegistry([DelimitedParser(), JsonParser(), HtmlParser()])

    with pytest.raises(UnsupportedDocumentError, match="UNSUPPORTED_FORMAT"):
        registry.resolve(source)


@pytest.mark.asyncio
async def test_docx_preserves_paragraph_table_order_without_fake_pages(tmp_path: Path) -> None:
    from docx import Document

    source = tmp_path / "report.docx"
    document = Document()
    document.add_heading("标题", level=1)
    document.add_paragraph("表前正文")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "金额"
    document.add_paragraph("表后正文")
    document.save(source)

    result = await DocxParser().parse(source, options=PipelineOptions())

    assert [segment.id for segment in result.segments] == [
        "D-P1",
        "D-P2",
        "D-T1",
        "D-P3",
    ]
    assert all(segment.location is None for segment in result.segments)
    assert isinstance(result.segments[0].source_locators[0], DocxBlockLocator)
    assert "| 项目 | 金额 |" in result.segments[2].text


@pytest.mark.asyncio
async def test_pptx_separates_slide_body_and_notes(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    source = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text = "正文内容"
    notes = slide.notes_slide.notes_text_frame
    notes.text = "备注内容"
    presentation.save(source)

    result = await PptxParser().parse(
        source,
        options=PipelineOptions(pptx_mode=PptxMode.TEXT_ONLY),
    )

    body = [segment for segment in result.segments if "正文内容" in segment.text]
    notes_segments = [segment for segment in result.segments if "备注内容" in segment.text]
    assert len(body) == len(notes_segments) == 1
    assert isinstance(body[0].source_locators[0], SlideObjectLocator)
    assert body[0].source_locators[0].content_kind == "body"
    assert notes_segments[0].source_locators[0].content_kind == "notes"
    assert result.total_units == result.covered_units == 1


@pytest.mark.asyncio
async def test_pptx_keeps_text_objects_without_geometry(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    source = tmp_path / "missing-geometry.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    positioned = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    positioned.text = "有坐标正文"
    inherited = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1))
    inherited.text = "无坐标正文"
    inherited._element.spPr.remove(inherited._element.spPr.xfrm)
    presentation.save(source)

    result = await PptxParser().parse(
        source,
        options=PipelineOptions(pptx_mode=PptxMode.TEXT_ONLY),
    )

    assert [segment.text for segment in result.segments] == ["有坐标正文", "无坐标正文"]
    assert result.parser_metadata["unpositioned_text_objects"] == 1
    assert any("1 个文本对象缺少坐标" in warning for warning in result.warnings)


@pytest.mark.asyncio
async def test_pptx_visual_transcript_covers_missing_native_geometry(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    source = tmp_path / "missing-geometry-visual.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    inherited = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    inherited.text = "无坐标正文"
    inherited._element.spPr.remove(inherited._element.spPr.xfrm)
    presentation.save(source)

    result = await PptxParser(
        vision=FakeSlideVision(),
        renderer=FakeSlideRenderer(),
    ).parse(source, options=PipelineOptions(pptx_mode=PptxMode.VISION_ALL))

    assert result.parser_metadata["unpositioned_text_objects"] == 1
    assert not any("文本对象缺少坐标" in warning for warning in result.warnings)
    assert result.warnings == []


@pytest.mark.asyncio
async def test_pptx_visual_primary_keeps_native_objects_as_auxiliary(
    tmp_path: Path,
) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    source = tmp_path / "visual-slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    first = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    first.text = "第一块正文"
    second = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1))
    second.text = "第二块正文"
    slide.notes_slide.notes_text_frame.text = "讲者备注"
    simple_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    simple = simple_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    simple.text = "单对象正文"
    presentation.save(source)
    renderer = FakeSlideRenderer()
    vision = FakeSlideVision()

    result = await PptxParser(vision=vision, renderer=renderer).parse(
        source,
        options=PipelineOptions(pptx_mode=PptxMode.AUTO),
    )

    assert renderer.slide_indices == (0, 1)
    assert vision.native_text[1] == "第一块正文\n\n第二块正文"
    visual = next(
        segment
        for segment in result.segments
        if segment.representation_id == "pptx:vision_transcript"
    )
    native = [
        segment for segment in result.segments if segment.representation_id == "pptx:native_objects"
    ]
    notes = next(segment for segment in result.segments if "讲者备注" in segment.text)
    assert visual.is_primary
    assert visual.source_locators == [
        SlideObjectLocator(
            slide_part="/ppt/slides/slide1.xml",
            slide_number=1,
            shape_id=None,
            content_kind="body",
        )
    ]
    assert native and not any(segment.is_primary for segment in native)
    assert notes.is_primary
    assert result.total_units == result.covered_units == 2
    assert result.parser_metadata["visual_slides_succeeded"] == 2
    assert result.warnings == []


@pytest.mark.asyncio
async def test_pptx_filters_large_single_letter_background_artifact(
    tmp_path: Path,
) -> None:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    source = tmp_path / "decorative.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(0),
        presentation.slide_width,
        presentation.slide_height,
    )
    background.text = "c"
    label = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    label.text = "A"
    presentation.save(source)

    result = await PptxParser().parse(
        source,
        options=PipelineOptions(pptx_mode=PptxMode.TEXT_ONLY),
    )

    assert [segment.text for segment in result.segments] == ["A"]
    assert result.parser_metadata["filtered_decorative_objects"] == 1


@pytest.mark.asyncio
async def test_xlsx_keeps_formulas_cached_values_and_structure_separate(
    tmp_path: Path,
) -> None:
    import openpyxl

    source = tmp_path / "预算.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "预算"
    sheet.append(["项目", "金额"])
    sheet.append(["总计", "=SUM(B3:B4)"])
    sheet.merge_cells("A3:A4")
    sheet["A3"] = "收入"
    sheet["B3"] = 10
    sheet["B4"] = 20
    hidden = workbook.create_sheet("隐藏表")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "保留"
    workbook.save(source)

    result = await XlsxParser().parse(source, options=PipelineOptions())

    formula_segment = next(segment for segment in result.segments if "formula=" in segment.text)
    locator = formula_segment.source_locators[0]
    assert isinstance(locator, CellRangeLocator)
    assert locator.sheet_name == "预算"
    assert "cached=<not cached>" in formula_segment.text
    budget_metadata = result.parser_metadata["sheets"][0]
    assert budget_metadata["merged_ranges"] == ["A3:A4"]
    assert any("隐藏工作表" in warning for warning in result.warnings)


def test_ooxml_extension_must_match_container_type(tmp_path: Path) -> None:
    import openpyxl

    source = tmp_path / "wrong.docx"
    workbook = openpyxl.Workbook()
    workbook.save(source)

    with pytest.raises(LocalDocumentError, match="文档类型不匹配"):
        DocxParser().validate(source)
