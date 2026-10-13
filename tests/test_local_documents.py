"""Local document parsing, visual routing, recap validation, and persistence."""

from __future__ import annotations

import asyncio
import re
import threading
import time
from pathlib import Path

import pytest

from searchos.local.__main__ import _parser as cli_parser
from searchos.local.models import (
    PdfMode,
    PipelineOptions,
    PptxMode,
    ProcessingStatus,
)
from searchos.local.parsers.image import ImageParser
from searchos.local.parsers.markdown import MarkdownParser
from searchos.local.parsers.pdf import PdfParser, _PageInput
from searchos.local.pipeline import DocumentRecapPipeline
from searchos.local.reader.references import parse_ref_id
from searchos.local.reader.tools import local_find, local_open, reset_local_browser


def _document_links(path: Path) -> list[tuple[str, str]]:
    return re.findall(r"^- \[([^]]+)]\((local://[^)]+)\)", path.read_text(encoding="utf-8"), re.M)


def test_remote_local_cli_uses_shared_model_configuration(monkeypatch) -> None:
    from searchos.config import setup_wizard, web_overlay
    from searchos.local import __main__ as local_cli

    calls: list[str] = []
    monkeypatch.setattr(
        setup_wizard,
        "ensure_model_config",
        lambda: calls.append("ensure"),
    )
    monkeypatch.setattr(
        web_overlay,
        "load_and_apply",
        lambda: calls.append("overlay"),
    )

    local_cli._configure_models(allow_remote=True)
    assert calls == ["ensure", "overlay"]

    calls.clear()
    local_cli._configure_models(allow_remote=False)
    assert calls == []


class FakeRecapClient:
    def __init__(self, *, bad_ids: bool = False) -> None:
        self.calls = 0
        self.bad_ids = bad_ids

    async def summarize_segments(self, segments):
        self.calls += 1
        if self.bad_ids:
            return {
                "sections": [
                    {
                        "title": segments[0].title,
                        "summary": "概括这一部分的主要内容。",
                        "segment_ids": ["unknown"],
                    }
                ]
            }
        return {
            "sections": [
                {
                    "title": segment.title,
                    "summary": "概括这一部分的主要内容。",
                    "segment_ids": [segment.id],
                }
                for segment in segments
            ]
        }

    async def summarize_document(self, sections):
        return "介绍文档的主题、范围和主要内容。"


@pytest.mark.asyncio
async def test_recap_progress_counts_completed_segments() -> None:
    from searchos.local.models import (
        DocumentKind,
        LineRangeLocator,
        ParsedDocument,
        ParsedSegment,
    )
    from searchos.local.recap.service import RecapService

    segments = [
        ParsedSegment(
            id=f"L{index}",
            title=f"第 {index} 节",
            text=str(index) * 1200,
            source_locators=[LineRangeLocator(start_line=index, end_line=index)],
        )
        for index in range(1, 4)
    ]
    document = ParsedDocument(
        document_id="progress",
        source_path="/tmp/progress.md",
        display_path="progress.md",
        name="progress.md",
        kind=DocumentKind.MARKDOWN,
        sha256="digest",
        modified_ns=0,
        size_bytes=3600,
        segments=segments,
        total_units=3,
        covered_units=3,
    )
    updates: list[tuple[int, int]] = []

    await RecapService(FakeRecapClient()).generate(
        document,
        max_chars=1500,
        on_progress=lambda completed, total: updates.append((completed, total)),
    )

    assert updates == [(0, 3), (1, 3), (2, 3), (3, 3)]


class FakeVision:
    def __init__(self, *, fail_pages: set[int] | None = None) -> None:
        self.pages: list[int] = []
        self.fail_pages = fail_pages or set()

    async def transcribe_page(self, *, page_number, image_data_uri, native_text):
        assert image_data_uri.startswith("data:image/jpeg;base64,")
        self.pages.append(page_number)
        if page_number in self.fail_pages:
            raise RuntimeError("vision unavailable")
        return f"视觉转录第 {page_number} 页"

    async def transcribe_slide(self, *, slide_number, image_data_uri, native_text):
        return await self.transcribe_page(
            page_number=slide_number,
            image_data_uri=image_data_uri,
            native_text=native_text,
        )


class MergingRecapClient(FakeRecapClient):
    async def summarize_segments(self, segments):
        self.calls += 1
        return {
            "sections": [
                {
                    "title": "合并标题",
                    "summary": "合并摘要。",
                    "segment_ids": [segment.id for segment in segments],
                }
            ]
        }


class CrossPageChapterRecapClient(FakeRecapClient):
    async def summarize_segments(self, segments):
        self.calls += 1
        assert [segment.id for segment in segments] == [
            "P1-native",
            "P2-native",
            "P3-native",
            "P4-native",
        ]
        return {
            "sections": [
                {
                    "title": "引言",
                    "summary": "介绍研究背景。",
                    "segment_ids": ["P1-native", "P2-native"],
                },
                {
                    "title": "方法",
                    "summary": "说明研究方法。",
                    "segment_ids": ["P2-native", "P3-native"],
                },
                {
                    "title": "结论",
                    "summary": "总结主要结论。",
                    "segment_ids": ["P4-native"],
                },
            ]
        }


class SingleSegmentProtocolRecapClient(FakeRecapClient):
    def __init__(self) -> None:
        super().__init__()
        self.single_calls = 0

    async def summarize_segments(self, segments):
        raise AssertionError("单个不可分片段不应让模型分配 segment_id")

    async def summarize_segment(self, segment):
        self.single_calls += 1
        return {"title": "模型标题", "summary": "概括单个片段的内容。"}


class FakeCatalogClient:
    async def classify_document(self, document, existing):
        if existing:
            return {
                "subindexes": [
                    {
                        "subindex_id": existing[0].subindex_id,
                        "title": existing[0].title,
                    }
                ]
            }
        return {"subindexes": [{"subindex_id": "", "title": "测试资料"}]}

    async def revise_subindex(self, title, summary, documents):
        return {
            "title": title,
            "summary": f"{title}包含 {len(documents)} 份文件。",
        }


class SlowCatalogClient(FakeCatalogClient):
    async def classify_document(self, document, existing):
        await asyncio.sleep(0.02)
        return await super().classify_document(document, existing)

    async def revise_subindex(self, title, summary, documents):
        await asyncio.sleep(0.02)
        return await super().revise_subindex(title, summary, documents)


class EmptyAssignmentCatalogClient(FakeCatalogClient):
    async def classify_document(self, document, existing):
        return {"subindexes": []}

    async def revise_subindex(self, title, summary, documents):
        return {
            "title": "保研与推免规划",
            "summary": "涵盖推免资格、准备与申请流程。",
        }


class BottomUpCatalogClient(FakeCatalogClient):
    async def classify_document(self, document, existing):
        if existing:
            return {
                "subindexes": [{
                    "subindex_id": existing[0].subindex_id,
                    "title": existing[0].title,
                }]
            }
        return {"subindexes": [{"subindex_id": "", "title": "升学资料"}]}

    async def revise_subindex(self, title, summary, documents):
        count = len(documents)
        return {
            "title": "推免政策" if count == 1 else "保研与推免规划",
            "summary": f"覆盖 {count} 份文档的推免资格、准备与申请流程。",
        }


class FakeRenderer:
    def render_page(self, path, *, page_index, dpi, max_dimension):
        assert Path(path).exists()
        assert dpi >= 72
        assert max_dimension >= 512
        return "data:image/jpeg;base64,dGVzdA=="


class FakeSlideRenderer:
    def render_slides(self, path, *, slide_indices, dpi, max_dimension):
        return {index: "data:image/jpeg;base64,dGVzdA==" for index in slide_indices}


class ConcurrentUnsafeRenderer(FakeRenderer):
    def __init__(self) -> None:
        self._lock = threading.Lock()
        self._active = 0
        self.max_active = 0

    def render_page(self, path, *, page_index, dpi, max_dimension):
        with self._lock:
            self._active += 1
            self.max_active = max(self.max_active, self._active)
        try:
            time.sleep(0.03)
            return super().render_page(
                path,
                page_index=page_index,
                dpi=dpi,
                max_dimension=max_dimension,
            )
        finally:
            with self._lock:
                self._active -= 1


@pytest.mark.asyncio
async def test_markdown_parser_preserves_headings_lines_and_fenced_hashes(tmp_path) -> None:
    source = tmp_path / "notes.md"
    source.write_text(
        "前言\n\n# 第一节\n内容\n```python\n# 不是标题\n```\n\n第二节\n------\n结尾\n",
        encoding="utf-8",
    )

    result = await MarkdownParser().parse(source, options=PipelineOptions())

    assert [(segment.title, segment.location.label) for segment in result.segments] == [
        ("文档开头", "L1-2"),
        ("第一节", "L3-8"),
        ("第二节", "L9-11"),
    ]
    assert "不是标题" in result.segments[1].text
    assert result.covered_units == result.total_units == 11


@pytest.mark.asyncio
async def test_pdf_auto_uses_vision_only_for_scan_or_complex_pages(tmp_path, monkeypatch) -> None:
    source = tmp_path / "mixed.pdf"
    source.write_bytes(b"fake")
    pages = [
        _PageInput(number=1, text="普通正文" * 100, image_count=0),
        _PageInput(number=2, text="", image_count=1),
        _PageInput(
            number=3,
            text="a    b    c    d\n1    2    3    4\n5    6    7    8\n9    10    11    12",
            image_count=0,
        ),
    ]
    monkeypatch.setattr("searchos.local.parsers.pdf._load_pdf_pages", lambda *a, **k: pages)
    vision = FakeVision()

    result = await PdfParser(vision=vision, renderer=FakeRenderer()).parse(
        source,
        options=PipelineOptions(pdf_native_min_chars=20),
    )

    assert vision.pages == [2, 3]
    assert [segment.extraction_method for segment in result.segments] == [
        "native_text",
        "vision_transcript",
        "native_text",
        "vision_transcript",
    ]
    assert [segment.is_primary for segment in result.segments] == [True, True, False, True]
    assert result.covered_units == result.total_units == 3


@pytest.mark.asyncio
async def test_pdf_keeps_native_and_vision_representations_selectable(
    tmp_path, monkeypatch
) -> None:
    source = tmp_path / "table.pdf"
    source.write_bytes(b"fake")
    native = "a    b    c    d\n1    2    3    4\n5    6    7    8\n9    10    11    12"
    monkeypatch.setattr(
        "searchos.local.parsers.pdf._load_pdf_pages",
        lambda *a, **k: [_PageInput(number=1, text=native, image_count=0)],
    )
    index = tmp_path / "index"
    result = await DocumentRecapPipeline(
        output_root=index,
        recap_client=FakeRecapClient(),
        vision_transcriber=FakeVision(),
        page_renderer=FakeRenderer(),
        recap_model_name="test",
        vision_model_name="test",
    ).process(
        source,
        document_id="pdf_representations",
        options=PipelineOptions(allow_remote=True, pdf_native_min_chars=20),
    )
    ref_id = _document_links(result.document_path)[0][1]
    reset_local_browser(index, source_roots=[tmp_path])

    primary = await local_open.ainvoke({"ref_id": ref_id})
    native_view = await local_open.ainvoke(
        {"ref_id": ref_id, "representation_id": "pdf:native_text"}
    )

    assert "视觉转录第 1 页" in primary
    assert native in native_view
    assert '"representation_id":"pdf:vision_transcript"' in primary
    assert '"representation_id":"pdf:native_text"' in native_view


@pytest.mark.asyncio
async def test_pptx_whole_slide_reference_can_open_native_auxiliary(
    tmp_path,
) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    source = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    first = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    first.text = "第一块"
    second = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1))
    second.text = "第二块"
    presentation.save(source)
    index = tmp_path / "index"
    result = await DocumentRecapPipeline(
        output_root=index,
        recap_client=FakeRecapClient(),
        vision_transcriber=FakeVision(),
        slide_renderer=FakeSlideRenderer(),
        recap_model_name="test",
        vision_model_name="test",
    ).process(
        source,
        document_id="pptx_representations",
        options=PipelineOptions(
            allow_remote=True,
            pptx_mode=PptxMode.VISION_ALL,
        ),
    )
    ref_id = _document_links(result.document_path)[0][1]
    reset_local_browser(index, source_roots=[tmp_path])

    primary = await local_open.ainvoke({"ref_id": ref_id})
    native = await local_open.ainvoke(
        {"ref_id": ref_id, "representation_id": "pptx:native_objects"}
    )

    assert "视觉转录第 1 页" in primary
    assert "第一块" in native and "第二块" in native
    assert '"representation_id":"pptx:vision_transcript"' in primary
    assert '"representation_id":"pptx:native_objects"' in native
    assert '"coverage":"partial"' in native


@pytest.mark.asyncio
async def test_image_without_vision_registers_source_without_invented_text(tmp_path) -> None:
    from PIL import Image

    source = tmp_path / "scan.png"
    Image.new("RGB", (8, 6), "white").save(source)

    result = await ImageParser().parse(source, options=PipelineOptions())

    assert result.total_units == 1
    assert result.covered_units == 0
    assert result.segments == []
    assert "不生成正文" in result.warnings[0]

    index = tmp_path / "index"
    processed = await DocumentRecapPipeline(output_root=index).process(
        source,
        document_id="image_without_text",
        options=PipelineOptions(allow_remote=False),
    )
    ref_id = _document_links(processed.document_path)[0][1]
    reset_local_browser(index, source_roots=[tmp_path])
    original = await local_open.ainvoke({"ref_id": ref_id, "view": "original"})

    assert isinstance(original, list)
    assert original[1]["type"] == "image"


@pytest.mark.asyncio
async def test_local_only_document_keeps_structural_navigation(tmp_path) -> None:
    source = tmp_path / "notes.md"
    source.write_text("# 标题\n正文\n", encoding="utf-8")
    index = tmp_path / "index"

    result = await DocumentRecapPipeline(output_root=index).process(
        source,
        document_id="local_only_notes",
        options=PipelineOptions(allow_remote=False),
    )
    links = _document_links(result.document_path)
    reset_local_browser(index, source_roots=[tmp_path])
    opened = await local_open.ainvoke({"ref_id": links[0][1]})

    assert links[0][0] == "L1-2"
    assert "正文" in opened


@pytest.mark.asyncio
async def test_image_vision_transcript_and_original_view_are_structured(tmp_path) -> None:
    from PIL import Image

    source = tmp_path / "scan.png"
    Image.new("RGB", (8, 6), "white").save(source)
    index = tmp_path / "index"
    result = await DocumentRecapPipeline(
        output_root=index,
        recap_client=FakeRecapClient(),
        vision_transcriber=FakeVision(),
        recap_model_name="test",
        vision_model_name="test",
    ).process(
        source,
        document_id="image_test",
        options=PipelineOptions(allow_remote=True),
    )
    ref_id = _document_links(result.document_path)[0][1]
    reset_local_browser(index, source_roots=[tmp_path])

    original = await local_open.ainvoke({"ref_id": ref_id, "view": "original"})

    assert isinstance(original, list)
    assert original[0]["type"] == "text"
    assert original[1]["type"] == "image"
    assert original[1]["mime_type"] == "image/jpeg"


@pytest.mark.asyncio
async def test_normalized_record_range_reads_all_backing_csv_blocks(tmp_path) -> None:
    source = tmp_path / "records.csv"
    source.write_text("name,value\na,1\nb,2\nc,3\n", encoding="utf-8")
    index = tmp_path / "index"
    result = await DocumentRecapPipeline(
        output_root=index,
        recap_client=MergingRecapClient(),
        recap_model_name="test",
        vision_model_name="test",
    ).process(
        source,
        document_id="csv_ranges",
        options=PipelineOptions(
            allow_remote=True,
            tabular_chunk_rows=1,
        ),
    )
    ref_id = _document_links(result.document_path)[0][1]
    reset_local_browser(index, source_roots=[tmp_path])

    opened = await local_open.ainvoke({"ref_id": ref_id})
    found = await local_find.ainvoke({"ref_id": ref_id, "pattern": "b,2"})

    assert all(value in opened for value in ("a,1", "b,2", "c,3"))
    assert '"coverage":"declared_blocks"' in opened
    assert "representation_id='primary'" in found


@pytest.mark.asyncio
async def test_concurrent_catalog_updates_do_not_lose_documents(tmp_path) -> None:
    first_source = tmp_path / "first.md"
    second_source = tmp_path / "second.md"
    first_source.write_text("# 第一份\n正文\n", encoding="utf-8")
    second_source.write_text("# 第二份\n正文\n", encoding="utf-8")
    index = tmp_path / "index"

    async def process(source: Path, document_id: str) -> None:
        await DocumentRecapPipeline(
            output_root=index,
            recap_client=FakeRecapClient(),
            catalog_client=SlowCatalogClient(),
            recap_model_name="test",
            vision_model_name="test",
        ).process(
            source,
            document_id=document_id,
            options=PipelineOptions(allow_remote=True, build_catalog=True),
        )

    await asyncio.gather(
        process(first_source, "first"),
        process(second_source, "second"),
    )

    subindex = next((index / "subindexes").glob("*.md")).read_text(encoding="utf-8")
    assert "documents/first.md" in subindex
    assert "documents/second.md" in subindex


@pytest.mark.asyncio
async def test_pdf_text_only_never_calls_vision(tmp_path, monkeypatch) -> None:
    source = tmp_path / "scan.pdf"
    source.write_bytes(b"fake")
    monkeypatch.setattr(
        "searchos.local.parsers.pdf._load_pdf_pages",
        lambda *a, **k: [_PageInput(number=1, text="", image_count=1)],
    )
    vision = FakeVision()

    result = await PdfParser(vision=vision, renderer=FakeRenderer()).parse(
        source,
        options=PipelineOptions(pdf_mode=PdfMode.TEXT_ONLY),
    )

    assert vision.pages == []
    assert result.covered_units == 0
    assert result.total_units == 1


@pytest.mark.asyncio
async def test_pdf_failed_scan_page_is_not_counted_as_covered(tmp_path, monkeypatch) -> None:
    source = tmp_path / "scan.pdf"
    source.write_bytes(b"fake")
    monkeypatch.setattr(
        "searchos.local.parsers.pdf._load_pdf_pages",
        lambda *a, **k: [_PageInput(number=1, text="", image_count=1)],
    )

    result = await PdfParser(
        vision=FakeVision(fail_pages={1}),
        renderer=FakeRenderer(),
    ).parse(source, options=PipelineOptions())

    assert result.covered_units == 0
    assert any("视觉转录失败" in warning for warning in result.warnings)


@pytest.mark.asyncio
async def test_pdf_renderer_calls_are_serialized(tmp_path, monkeypatch) -> None:
    source = tmp_path / "visual.pdf"
    source.write_bytes(b"fake")
    monkeypatch.setattr(
        "searchos.local.parsers.pdf._load_pdf_pages",
        lambda *a, **k: [
            _PageInput(number=number, text=f"第 {number} 页正文", image_count=1)
            for number in range(1, 5)
        ],
    )
    renderer = ConcurrentUnsafeRenderer()

    result = await PdfParser(vision=FakeVision(), renderer=renderer).parse(
        source,
        options=PipelineOptions(
            pdf_mode=PdfMode.VISION_ALL,
            pdf_vision_concurrency=4,
        ),
    )

    assert result.covered_units == 4
    assert renderer.max_active == 1


@pytest.mark.asyncio
async def test_pdf_recap_groups_pages_into_chapters_and_allows_boundary_overlap(
    tmp_path, monkeypatch
) -> None:
    source = tmp_path / "paper.pdf"
    source.write_bytes(b"fake")
    monkeypatch.setattr(
        "searchos.local.parsers.pdf._load_pdf_pages",
        lambda *a, **k: [
            _PageInput(number=number, text="正文" * 100, image_count=0) for number in range(1, 5)
        ],
    )
    client = CrossPageChapterRecapClient()
    pipeline = DocumentRecapPipeline(
        output_root=tmp_path / "index",
        recap_client=client,
        recap_model_name="qwen3.6-35b-a3b",
        vision_model_name="qwen3.6-35b-a3b",
    )

    result = await pipeline.process(
        source,
        document_id="pdf_test",
        options=PipelineOptions(
            allow_remote=True,
            pdf_mode=PdfMode.TEXT_ONLY,
        ),
    )

    assert result.status == ProcessingStatus.READY
    assert client.calls == 1
    links = _document_links(result.document_path)
    assert [label for label, _ in links] == ["P1-2", "P2-3", "P4"]
    assert all(parse_ref_id(ref_id).document_id == "pdf_test" for _, ref_id in links)
    rendered = result.document_path.read_text(encoding="utf-8")
    assert "介绍研究背景。 (400 chars)" in rendered
    assert "说明研究方法。 (400 chars)" in rendered


@pytest.mark.asyncio
async def test_pdf_removes_binary_control_characters(tmp_path, monkeypatch) -> None:
    source = tmp_path / "text.pdf"
    source.write_bytes(b"fake")
    monkeypatch.setattr(
        "searchos.local.parsers.pdf._load_pdf_pages",
        lambda *a, **k: [_PageInput(number=1, text="alpha\x00beta\ntext", image_count=0)],
    )

    result = await PdfParser().parse(
        source,
        options=PipelineOptions(pdf_mode=PdfMode.TEXT_ONLY),
    )

    assert result.segments[0].text == "alphabeta\ntext"


@pytest.mark.asyncio
async def test_pipeline_writes_minimal_chapter_index_and_reuses_cache(tmp_path) -> None:
    source_root = tmp_path / "source"
    source_root.mkdir()
    source = source_root / "设计.md"
    source.write_text("# 范围\n正文\n\n# 方法\n实现细节\n", encoding="utf-8")
    output = tmp_path / "index"
    client = FakeRecapClient()
    pipeline = DocumentRecapPipeline(
        output_root=output,
        recap_client=client,
        recap_model_name="qwen3.6-35b-a3b",
        vision_model_name="qwen3.6-35b-a3b",
    )
    options = PipelineOptions(allow_remote=True)

    first = await pipeline.process(
        source,
        source_root=source_root,
        document_id="doc_test",
        options=options,
    )

    assert first.status == ProcessingStatus.READY
    links = _document_links(first.document_path)
    assert [label for label, _ in links] == ["L1-3", "L4-5"]
    assert all(parse_ref_id(ref_id).document_id == "doc_test" for _, ref_id in links)
    assert 'extraction_method":"native"' in first.text_path.read_text(encoding="utf-8")

    cached = await pipeline.process(
        source,
        source_root=source_root,
        document_id="doc_test",
        options=options,
    )
    assert cached.cached
    assert client.calls == 1


@pytest.mark.asyncio
async def test_single_segment_recap_keeps_program_owned_location(tmp_path) -> None:
    source = tmp_path / "notes.md"
    source.write_text("# 原文标题\n正文\n", encoding="utf-8")
    client = SingleSegmentProtocolRecapClient()

    result = await DocumentRecapPipeline(
        output_root=tmp_path / "index",
        recap_client=client,
        recap_model_name="test",
        vision_model_name="test",
    ).process(
        source,
        document_id="single_segment",
        options=PipelineOptions(allow_remote=True, build_catalog=False),
    )

    assert result.status == ProcessingStatus.READY
    assert client.single_calls == 1
    assert result.recap is not None
    assert result.recap.sections[0].title == "原文标题"
    assert result.recap.sections[0].segment_ids == ["L1-2"]


def test_cli_generates_remote_recap_by_default_and_supports_local_only() -> None:
    parser = cli_parser()

    assert parser.parse_args(["notes.md"]).allow_remote is True
    assert parser.parse_args(["notes.md"]).build_catalog is True
    assert parser.parse_args(["slides.pptx"]).pptx_mode == "auto"
    assert parser.parse_args(["slides.pptx", "--pptx-mode", "vision_all"]).pptx_mode == "vision_all"
    assert parser.parse_args(["notes.md", "--local-only"]).allow_remote is False


@pytest.mark.asyncio
async def test_local_only_pipeline_does_not_resolve_model_profiles(tmp_path, monkeypatch) -> None:
    source = tmp_path / "notes.md"
    source.write_text("# 标题\n正文\n", encoding="utf-8")

    def fail_profile_lookup(role: str) -> str:
        raise AssertionError(f"不应解析模型角色：{role}")

    monkeypatch.setattr("searchos.local.pipeline._profile_model", fail_profile_lookup)
    result = await DocumentRecapPipeline(output_root=tmp_path / "index").process(
        source,
        options=PipelineOptions(allow_remote=False, build_catalog=False),
    )

    assert result.status == ProcessingStatus.PARTIAL
    assert result.text_path.exists()


@pytest.mark.asyncio
async def test_pipeline_builds_index_and_subindex_bottom_up(tmp_path) -> None:
    source_root = tmp_path / "source"
    source_root.mkdir()
    first_source = source_root / "a.md"
    second_source = source_root / "b.md"
    first_source.write_text("# A\n第一份正文\n", encoding="utf-8")
    second_source.write_text("# B\n第二份正文\n", encoding="utf-8")
    output = tmp_path / "index"
    pipeline = DocumentRecapPipeline(
        output_root=output,
        recap_client=FakeRecapClient(),
        catalog_client=FakeCatalogClient(),
        recap_model_name="qwen3.6-35b-a3b",
        vision_model_name="qwen3.6-35b-a3b",
    )
    options = PipelineOptions(allow_remote=True, build_catalog=True)

    await pipeline.process(
        first_source,
        source_root=source_root,
        document_id="doc_a",
        options=options,
    )
    result = await pipeline.process(
        second_source,
        source_root=source_root,
        document_id="doc_b",
        options=options,
    )

    index = (output / "INDEX.md").read_text(encoding="utf-8")
    subindex_path = next((output / "subindexes").glob("*.md"))
    subindex = subindex_path.read_text(encoding="utf-8")
    assert subindex_path.name == "测试资料.md"
    assert "[测试资料](subindexes/" in index
    assert "测试资料包含 2 份文件。" in index
    assert "[a.md](../documents/doc_a.md)" in subindex
    assert "[b.md](../documents/doc_b.md)" in subindex
    assert result.catalog_paths[0] == output / "INDEX.md"


@pytest.mark.asyncio
async def test_catalog_preserves_topic_identity_and_updates_recap_bottom_up(tmp_path) -> None:
    source_root = tmp_path / "source"
    source_root.mkdir()
    first_source = source_root / "policy.md"
    second_source = source_root / "timeline.md"
    first_source.write_text("# 推免政策\n资格与成绩要求。\n", encoding="utf-8")
    second_source.write_text("# 申请规划\n夏令营与预推免时间线。\n", encoding="utf-8")
    output = tmp_path / "index"
    pipeline = DocumentRecapPipeline(
        output_root=output,
        recap_client=FakeRecapClient(),
        catalog_client=BottomUpCatalogClient(),
        recap_model_name="test-model",
        vision_model_name="test-model",
    )
    options = PipelineOptions(allow_remote=True, build_catalog=True)

    await pipeline.process(
        first_source,
        source_root=source_root,
        document_id="doc_policy",
        options=options,
    )
    assert (output / "subindexes" / "升学资料.md").is_file()

    await pipeline.process(
        second_source,
        source_root=source_root,
        document_id="doc_timeline",
        options=options,
    )

    renamed = output / "subindexes" / "升学资料.md"
    assert renamed.is_file()
    assert "doc_policy.md" in renamed.read_text(encoding="utf-8")
    assert "doc_timeline.md" in renamed.read_text(encoding="utf-8")
    index = (output / "INDEX.md").read_text(encoding="utf-8")
    assert "[升学资料](subindexes/升学资料.md)" in index
    assert "覆盖 2 份文档的推免资格、准备与申请流程。" in index


@pytest.mark.asyncio
async def test_pipeline_revises_provisional_topic_when_catalog_returns_no_assignment(
    tmp_path,
) -> None:
    source = tmp_path / "recommendation.md"
    source.write_text("# 保研规划\n资格、夏令营与预推免。\n", encoding="utf-8")
    output = tmp_path / "index"
    pipeline = DocumentRecapPipeline(
        output_root=output,
        recap_client=FakeRecapClient(),
        catalog_client=EmptyAssignmentCatalogClient(),
        recap_model_name="test-model",
        vision_model_name="test-model",
    )

    result = await pipeline.process(
        source,
        document_id="doc_uncategorized",
        options=PipelineOptions(allow_remote=True, build_catalog=True),
    )

    assert result.status == ProcessingStatus.READY
    assert (output / "subindexes" / "未分类.md").exists()
    assert not (output / "subindexes" / "保研与推免规划.md").exists()
    assert "[未分类](subindexes/未分类.md)" in (
        output / "INDEX.md"
    ).read_text(encoding="utf-8")


def test_catalog_preserves_stored_topic_id(tmp_path) -> None:
    from searchos.local.index.catalog import CatalogMarkdownStore

    root = tmp_path / "index"
    subindexes = root / "subindexes"
    subindexes.mkdir(parents=True)
    (root / "INDEX.md").write_text(
        "# Local Index\n\n- [测试资料](subindexes/sub_deadbeef1234.md)：测试摘要。 (12 chars)\n",
        encoding="utf-8",
    )
    (subindexes / "sub_deadbeef1234.md").write_text(
        "# 测试资料\n\n- [a.md](../documents/doc_a.md)：文件摘要。 (12 chars)\n",
        encoding="utf-8",
    )
    store = CatalogMarkdownStore(root)

    store.write(store.load())

    assert not (subindexes / "测试资料.md").exists()
    assert (subindexes / "sub_deadbeef1234.md").is_file()
    assert "subindexes/sub_deadbeef1234.md" in (root / "INDEX.md").read_text(encoding="utf-8")


@pytest.mark.asyncio
async def test_invalid_model_locations_leave_recap_partial(tmp_path) -> None:
    source = tmp_path / "notes.md"
    source.write_text("# 标题\n正文\n", encoding="utf-8")
    pipeline = DocumentRecapPipeline(
        output_root=tmp_path / "index",
        recap_client=FakeRecapClient(bad_ids=True),
        recap_model_name="qwen3.6-35b-a3b",
        vision_model_name="qwen3.6-35b-a3b",
    )

    result = await pipeline.process(
        source,
        document_id="notes_test",
        options=PipelineOptions(allow_remote=True),
    )

    assert result.status == ProcessingStatus.PARTIAL
    assert result.recap is not None and not result.recap.sections
    assert any("未知 segment_id" in warning for warning in result.warnings)

    retried = await pipeline.process(
        source,
        options=PipelineOptions(allow_remote=True),
    )
    assert not retried.cached


@pytest.mark.asyncio
async def test_enabling_remote_recap_invalidates_parse_only_cache(tmp_path) -> None:
    source = tmp_path / "notes.md"
    source.write_text("# 标题\n正文\n", encoding="utf-8")
    client = FakeRecapClient()
    pipeline = DocumentRecapPipeline(
        output_root=tmp_path / "index",
        recap_client=client,
        recap_model_name="qwen3.6-35b-a3b",
        vision_model_name="qwen3.6-35b-a3b",
    )

    parsed_only = await pipeline.process(source, options=PipelineOptions())
    with_recap = await pipeline.process(
        source,
        options=PipelineOptions(allow_remote=True),
    )

    assert parsed_only.status == ProcessingStatus.PARTIAL
    assert not with_recap.cached
    assert with_recap.status == ProcessingStatus.READY


@pytest.mark.asyncio
async def test_model_initialization_failure_still_persists_extracted_text(
    tmp_path, monkeypatch
) -> None:
    source = tmp_path / "notes.md"
    source.write_text("# 标题\n正文\n", encoding="utf-8")

    def fail_model():
        raise RuntimeError("missing key")

    monkeypatch.setattr("searchos.local.pipeline.ModelRecapClient", fail_model)
    pipeline = DocumentRecapPipeline(
        output_root=tmp_path / "index",
        recap_model_name="qwen3.6-35b-a3b",
        vision_model_name="qwen3.6-35b-a3b",
    )

    result = await pipeline.process(
        source,
        document_id="notes_test",
        options=PipelineOptions(allow_remote=True),
    )

    assert result.status == ProcessingStatus.PARTIAL
    assert result.text_path.exists()
    assert "正文" in result.text_path.read_text(encoding="utf-8")
    assert any("模型初始化失败" in warning for warning in result.warnings)


@pytest.mark.asyncio
async def test_explicit_markdown_headings_are_not_merged(tmp_path) -> None:
    source = tmp_path / "notes.md"
    source.write_text("# 第一节\n正文\n# 第二节\n正文\n", encoding="utf-8")
    client = MergingRecapClient()
    pipeline = DocumentRecapPipeline(
        output_root=tmp_path / "index",
        recap_client=client,
        recap_model_name="qwen3.6-35b-a3b",
        vision_model_name="qwen3.6-35b-a3b",
    )

    result = await pipeline.process(
        source,
        document_id="notes_test",
        options=PipelineOptions(allow_remote=True),
    )

    assert result.status == ProcessingStatus.READY
    assert result.recap is not None
    assert [section.title for section in result.recap.sections] == ["第一节", "第二节"]
    assert client.calls == 3
    links = _document_links(result.document_path)
    assert [label for label, _ in links] == ["L1-2", "L3-4"]
    assert all(parse_ref_id(ref_id).document_id == "notes_test" for _, ref_id in links)
