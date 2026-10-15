"""短文档一次读取多页，保留版本、范围和证据定位约束。"""

from __future__ import annotations

import hashlib
import re
from types import SimpleNamespace

import pytest
import pytest_asyncio

from searchos.harness.middleware.extraction import ExtractionReceipt, LocalEvidenceMiddleware
from searchos.harness.middleware.models import SourceContext
from searchos.local.index.source_registry import SourceRegistry
from searchos.local.models import PipelineOptions
from searchos.local.pipeline import DocumentRecapPipeline
from searchos.local.reader.references import parse_ref_id, resolve_reference
from searchos.local.reader.tools import local_catalog, local_open, selected_document_scope


@pytest_asyncio.fixture
async def deck(tmp_path):
    from pptx import Presentation

    presentation = Presentation()
    for number in range(1, 10):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"Study {number}"
        slide.placeholders[1].text = f"Finding {number}: the source supports this result."
    source = tmp_path / "study.pptx"
    presentation.save(source)
    root = tmp_path / "index"
    result = await DocumentRecapPipeline(output_root=root).process(
        source, options=PipelineOptions(
            allow_remote=False, build_catalog=False,
        ),
    )
    return root, source, result, SourceRegistry(root).read(result.document_id)


async def document_ref(record):
    catalog = await local_catalog.ainvoke({"path": f"documents/{record.document_id}.md"})
    refs = re.findall(r'ref_id="(local://[^\"]+)"', catalog)
    whole = [ref for ref in refs if "/all?" in ref]
    assert len(whole) == 1, "目录应提供一次读取已索引全文的入口"
    assert "[LOCAL_SOURCE]" not in catalog
    return whole[0], [ref for ref in refs if ref != whole[0]]


@pytest.mark.asyncio
async def test_nine_slide_reading_uses_one_source_window_and_preserves_evidence(deck):
    root, _, _, record = deck
    with selected_document_scope(root, [record]):
        ref_id, sections = await document_ref(record)
        assert len(sections) == len(record.references) >= 9
        message = await local_open.ainvoke({
            "type": "tool_call", "name": "local_open", "id": "read-deck",
            "args": {"ref_id": ref_id, "max_chars": 20000},
        })
        assert "[END OF SOURCE]" in message.content
        body = message.content.split("\n\n", 1)[1]
        for number in range(1, 10):
            assert body.count(f"Finding {number}:") == 1
        metadata = message.artifact["local_source_metadata"]
        assert {loc["slide_number"] for loc in metadata["source_locator"]} == set(range(1, 10))
        assert metadata["source_revision"] == record.source_revision
        assert metadata["excerpt_hash"] == "sha256:" + hashlib.sha256(body.encode()).hexdigest()

        observations = []

        async def submit(observation, **kwargs):
            observations.append(observation)
            return ExtractionReceipt(accepted=True)

        middleware = LocalEvidenceMiddleware(
            SimpleNamespace(submit=submit),
            source=SourceContext(scope={"documents": [record.model_dump(mode="json")]}),
        )

        async def read(_request):
            return message

        await middleware.awrap_tool_call(SimpleNamespace(tool_call={"name": "local_open"}), read)
        assert len(observations) == 1
        canonical = observations[0].source_url
        replayed = await local_open.ainvoke({"ref_id": canonical, "max_chars": 20000})
        assert replayed.split("\n\n", 1)[1] == body


@pytest.mark.asyncio
async def test_document_windows_continue_without_dropping_or_repeating_text(deck):
    root, _, _, record = deck
    with selected_document_scope(root, [record]):
        ref_id, _ = await document_ref(record)
        full = await local_open.ainvoke({"ref_id": ref_id, "max_chars": 20000})
        boundary = full.split("\n\n", 1)[1].index("supports") + 3
        first = await local_open.ainvoke({"ref_id": ref_id, "max_chars": boundary})
        rest = await local_open.ainvoke({"ref_id": ref_id, "loc": boundary, "max_chars": 20000})
        assert f'loc={boundary}, max_chars={boundary}' in first
        assert first.split("\n\n", 1)[1] + rest.split("\n\n", 1)[1] == full.split("\n\n", 1)[1]


@pytest.mark.asyncio
async def test_document_reference_obeys_scope_version_and_integrity(deck):
    root, source, result, record = deck
    with selected_document_scope(root, [record]):
        ref_id, _ = await document_ref(record)
    with selected_document_scope(root, []):
        assert "选中文档范围" in await local_open.ainvoke({"ref_id": ref_id})
    with selected_document_scope(root, [record]):
        original = source.read_bytes()
        source.write_bytes(original + b"changed")
        assert "VERSION_MISMATCH" in await local_open.ainvoke({"ref_id": ref_id})
        source.write_bytes(original)
        result.text_path.write_text(result.text_path.read_text() + "changed")
        assert "CORRUPT_ARTIFACT" in await local_open.ainvoke({"ref_id": ref_id})


def test_document_reference_resolves_only_declared_locations():
    from searchos.local.models import LineRangeLocator

    ref = parse_ref_id("local://doc_test/all?v=" + "a" * 24)
    ranges = [LineRangeLocator(start_line=1, end_line=2),
              LineRangeLocator(start_line=8, end_line=9)]
    resolved = resolve_reference(ref, {"s1": ranges[:1], "s2": ranges[1:]},
                                 library_id="lib_test", source_revision="a" * 64)
    assert list(resolved.locators) == ranges


def test_document_reference_keeps_numeric_section_order_after_registry_serialization():
    from searchos.local.models import LineRangeLocator

    ref = parse_ref_id("local://doc_test/all?v=" + "a" * 24)
    ranges = [LineRangeLocator(start_line=n * 10, end_line=n * 10 + 1)
              for n in range(1, 12)]
    # YAML stores keys lexically: s1, s10, s11, s2, ... .
    references = dict(sorted((f"s{n}", [loc]) for n, loc in enumerate(ranges, 1)))
    resolved = resolve_reference(ref, references, library_id="lib_test",
                                 source_revision="a" * 64)
    assert list(resolved.locators) == ranges


def test_document_reference_preserves_reordered_ranges():
    from searchos.local.models import LineRangeLocator
    from searchos.local.reader.references import build_ref_id

    ref = parse_ref_id("local://doc_test/all?v=" + "a" * 24)
    ranges = [LineRangeLocator(start_line=8, end_line=9),
              LineRangeLocator(start_line=1, end_line=2)]
    resolved = resolve_reference(ref, {"s1": ranges[:1], "s2": ranges[1:]},
                                 library_id="lib_test", source_revision="a" * 64)
    canonical = build_ref_id(library_id="lib_test", document_id="doc_test",
                             source_revision="a" * 64, locators=list(resolved.locators))
    assert list(parse_ref_id(canonical).locators) == ranges


@pytest.mark.asyncio
async def test_oversized_locator_sets_keep_section_navigation(deck):
    from searchos.local.models import SlideObjectLocator

    root, _, _, record = deck
    record.references = {
        f"s{n}": [SlideObjectLocator(slide_part=f"/ppt/slides/slide{n}.xml", slide_number=n)]
        for n in range(1, 201)
    }
    ref_id = f"local://{record.document_id}/all?v={record.source_revision[:24]}"
    with selected_document_scope(root, [record]):
        catalog = await local_catalog.ainvoke({"path": f"documents/{record.document_id}.md"})
        assert "/all?" not in catalog
        assert "/s1?" in catalog
        assert "章节引用" in await local_open.ainvoke({"ref_id": ref_id})
