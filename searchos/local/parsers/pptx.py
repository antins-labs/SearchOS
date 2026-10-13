"""PPTX object extraction with optional whole-slide visual transcription."""

from __future__ import annotations

import asyncio
import re
from dataclasses import dataclass
from pathlib import Path

from searchos.local.models import (
    DocumentKind,
    ParsedSegment,
    ParseResult,
    PipelineOptions,
    PptxMode,
    SlideObjectLocator,
)
from searchos.local.errors import MissingDependencyError
from searchos.local.parsers.base import ParserCapability
from searchos.local.parsers.ooxml import inspect_ooxml_limits, validate_ooxml
from searchos.local.parsers.vision import (
    LibreOfficeSlideRenderer,
    SlideRenderer,
    SlideVisionTranscriber,
)

_PPTX_CONTENT_TYPE = (
    b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
)


def _table_text(table) -> str:
    return "\n".join(
        "\t".join(cell.text.replace("\t", " ").replace("\n", " ") for cell in row.cells)
        for row in table.rows
    )


def _shape_text(shape) -> str:
    if getattr(shape, "has_table", False):
        return _table_text(shape.table).strip()
    if getattr(shape, "has_text_frame", False):
        return "\n".join(
            paragraph.text.strip()
            for paragraph in shape.text_frame.paragraphs
            if paragraph.text.strip()
        ).strip()
    return ""


def _is_decorative_text(shape, text: str, *, auto_shape_type, slide_area: int) -> bool:
    compact = "".join(text.split())
    if not re.fullmatch(r"[A-Za-z]", compact):
        return False
    if shape.shape_type != auto_shape_type or slide_area <= 0:
        return False
    if shape.width is None or shape.height is None:
        return False
    return int(shape.width) * int(shape.height) >= slide_area // 5


def _shape_position(
    shape,
    *,
    slide_width: int,
    slide_height: int,
) -> tuple[int, int, bool]:
    """Return a stable sort position even when OOXML inherits no geometry."""

    has_position = shape.top is not None and shape.left is not None
    top = int(shape.top) if shape.top is not None else slide_height + 1
    left = int(shape.left) if shape.left is not None else slide_width + 1
    return top, left, has_position


@dataclass(frozen=True, slots=True)
class _SlideData:
    number: int
    part: str
    body: tuple[ParsedSegment, ...]
    notes: tuple[ParsedSegment, ...]
    native_text: str
    visual_objects: int
    group_objects: int

    @property
    def needs_ordering(self) -> bool:
        return len(self.body) > 1


class PptxParser:
    capability = ParserCapability(
        kind=DocumentKind.PPTX,
        extensions=(".pptx",),
        locator_kinds=("slide_object",),
        dependencies=("python-pptx", "LibreOffice (视觉模式)"),
        visual=True,
    )

    def __init__(
        self,
        *,
        vision: SlideVisionTranscriber | None = None,
        renderer: SlideRenderer | None = None,
    ) -> None:
        self._vision = vision
        self._renderer = renderer or LibreOfficeSlideRenderer()

    def validate(self, path: Path) -> None:
        validate_ooxml(path, content_type_marker=_PPTX_CONTENT_TYPE)

    async def parse(self, path: Path, *, options: PipelineOptions) -> ParseResult:
        try:
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE
        except ImportError as exc:
            raise MissingDependencyError(
                "PPTX 处理需要 local-docs 依赖；请安装 searchos[local-docs]"
            ) from exc

        inspect_ooxml_limits(
            path,
            max_uncompressed_bytes=options.max_archive_uncompressed_bytes,
            max_members=options.max_office_objects,
        )
        presentation = Presentation(str(path))
        slide_width = int(presentation.slide_width or 0)
        slide_height = int(presentation.slide_height or 0)
        slide_area = slide_width * slide_height
        visual_types = {
            MSO_SHAPE_TYPE.PICTURE,
            MSO_SHAPE_TYPE.CHART,
            getattr(MSO_SHAPE_TYPE, "MEDIA", None),
        }
        slides: list[_SlideData] = []
        object_count = 0
        filtered_decorations = 0
        unpositioned_text_objects = 0
        unpositioned_by_slide: dict[int, int] = {}

        for slide_number, slide in enumerate(presentation.slides, start=1):
            slide_part = str(slide.part.partname)
            positioned_body: list[tuple[int, int, ParsedSegment]] = []
            visual_objects = 0
            group_objects = 0
            for shape in slide.shapes:
                object_count += 1
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    group_objects += 1
                if shape.shape_type in visual_types:
                    visual_objects += 1
                text = _shape_text(shape)
                if not text:
                    continue
                if _is_decorative_text(
                    shape,
                    text,
                    auto_shape_type=MSO_SHAPE_TYPE.AUTO_SHAPE,
                    slide_area=slide_area,
                ):
                    filtered_decorations += 1
                    continue
                top, left, has_position = _shape_position(
                    shape,
                    slide_width=slide_width,
                    slide_height=slide_height,
                )
                if not has_position:
                    unpositioned_text_objects += 1
                    unpositioned_by_slide[slide_number] = (
                        unpositioned_by_slide.get(slide_number, 0) + 1
                    )
                positioned_body.append(
                    (
                        top,
                        left,
                        ParsedSegment(
                            id=f"S{slide_number}-O{shape.shape_id}",
                            title=f"幻灯片 {slide_number} 对象 {shape.shape_id}",
                            text=text,
                            source_locators=[
                                SlideObjectLocator(
                                    slide_part=slide_part,
                                    slide_number=slide_number,
                                    shape_id=shape.shape_id,
                                    content_kind="body",
                                )
                            ],
                            extraction_method="native_text",
                            representation_id="pptx:native_objects",
                        ),
                    )
                )

            positioned_body.sort(key=lambda item: (item[0], item[1], item[2].id))
            body = tuple(item[2] for item in positioned_body)
            notes: list[ParsedSegment] = []
            if slide.has_notes_slide:
                for shape in slide.notes_slide.shapes:
                    object_count += 1
                    text = _shape_text(shape)
                    if not text:
                        continue
                    notes.append(
                        ParsedSegment(
                            id=f"S{slide_number}-N{shape.shape_id}",
                            title=f"幻灯片 {slide_number} 备注",
                            text=text,
                            source_locators=[
                                SlideObjectLocator(
                                    slide_part=str(slide.notes_slide.part.partname),
                                    slide_number=slide_number,
                                    shape_id=shape.shape_id,
                                    content_kind="notes",
                                )
                            ],
                            extraction_method="native_text",
                            representation_id="pptx:native_notes",
                        )
                    )
            if object_count > options.max_office_objects:
                raise ValueError(f"PPTX 对象数超过限制 {options.max_office_objects}")
            slides.append(
                _SlideData(
                    number=slide_number,
                    part=slide_part,
                    body=body,
                    notes=tuple(notes),
                    native_text="\n\n".join(segment.text for segment in body),
                    visual_objects=visual_objects,
                    group_objects=group_objects,
                )
            )

        if options.pptx_mode == PptxMode.VISION_ALL:
            requested = {slide.number for slide in slides}
        elif options.pptx_mode == PptxMode.AUTO:
            requested = (
                {slide.number for slide in slides}
                if self._vision is not None
                else {
                    slide.number
                    for slide in slides
                    if slide.visual_objects or slide.group_objects or slide.needs_ordering
                }
            )
        else:
            requested = set()

        warnings: list[str] = []
        rendered: dict[int, str] = {}
        if requested and self._vision is None:
            warnings.append(f"{len(requested)} 张幻灯片需要视觉解析，但视觉模型未启用")
        elif requested:
            try:
                rendered = await asyncio.to_thread(
                    self._renderer.render_slides,
                    path,
                    slide_indices=tuple(number - 1 for number in sorted(requested)),
                    dpi=options.pptx_render_dpi,
                    max_dimension=options.pptx_max_dimension,
                )
            except Exception as exc:  # noqa: BLE001
                warnings.append(f"PPTX 整页渲染失败：{exc}")

        semaphore = asyncio.Semaphore(options.pptx_vision_concurrency)

        async def transcribe(slide: _SlideData) -> tuple[int, str, str]:
            image_data_uri = rendered.get(slide.number - 1)
            if image_data_uri is None or self._vision is None:
                return slide.number, "", ""
            try:
                async with semaphore:
                    text = await self._vision.transcribe_slide(
                        slide_number=slide.number,
                        image_data_uri=image_data_uri,
                        native_text=slide.native_text,
                    )
                if not text.strip():
                    return slide.number, "", "视觉转录返回空内容"
                return slide.number, text.strip(), ""
            except Exception as exc:  # noqa: BLE001
                return slide.number, "", str(exc)

        visual_results = await asyncio.gather(
            *(transcribe(slide) for slide in slides if slide.number in requested)
        )
        visual_text = {number: text for number, text, _ in visual_results if text}
        for number, _, error in visual_results:
            if error:
                warnings.append(f"Slide {number} 视觉转录失败：{error}")

        segments: list[ParsedSegment] = []
        for slide in slides:
            transcript = visual_text.get(slide.number)
            if transcript:
                segments.append(
                    ParsedSegment(
                        id=f"S{slide.number}-vision",
                        title=f"幻灯片 {slide.number}",
                        text=transcript,
                        source_locators=[
                            SlideObjectLocator(
                                slide_part=slide.part,
                                slide_number=slide.number,
                                shape_id=None,
                                content_kind="body",
                            )
                        ],
                        extraction_method="vision_transcript",
                        representation_id="pptx:vision_transcript",
                        is_primary=True,
                    )
                )
            segments.extend(
                segment.model_copy(update={"is_primary": not bool(transcript)})
                for segment in slide.body
            )
            segments.extend(slide.notes)

        visual_uncovered = sum(
            slide.visual_objects for slide in slides if slide.number not in visual_text
        )
        unpositioned_uncovered = sum(
            count
            for slide_number, count in unpositioned_by_slide.items()
            if slide_number not in visual_text
        )
        groups_uncovered = sum(
            slide.group_objects for slide in slides if slide.number not in visual_text
        )
        unordered = sum(
            1 for slide in slides if slide.needs_ordering and slide.number not in visual_text
        )
        if unordered:
            warnings.append(f"{unordered} 张幻灯片按对象位置近似排序；不保证视觉阅读顺序")
        if unpositioned_uncovered:
            warnings.append(
                f"{unpositioned_uncovered} 个文本对象缺少坐标；已保留并排在页面末尾"
            )
        if visual_uncovered:
            warnings.append(f"{visual_uncovered} 个图片、图表或媒体对象尚未视觉解析")
        if groups_uncovered:
            warnings.append(f"{groups_uncovered} 个组合对象尚未视觉解析")

        covered_slides = {
            slide.number
            for slide in slides
            if slide.number in visual_text
            or (slide.visual_objects == 0 and slide.group_objects == 0)
        }
        return ParseResult(
            kind=DocumentKind.PPTX,
            segments=segments,
            total_units=len(slides),
            covered_units=len(covered_slides),
            page_count=len(slides),
            warnings=warnings,
            parser_metadata={
                "slides": len(slides),
                "native_text_objects": sum(len(slide.body) for slide in slides),
                "visual_objects_uncovered": visual_uncovered,
                "group_objects_uncovered": groups_uncovered,
                "visual_slides_requested": len(requested),
                "visual_slides_succeeded": len(visual_text),
                "filtered_decorative_objects": filtered_decorations,
                "unpositioned_text_objects": unpositioned_text_objects,
                "unpositioned_text_objects_uncovered": unpositioned_uncovered,
                "ordering": (
                    "whole_slide_visual_primary_with_native_object_auxiliary"
                    if visual_text
                    else "top_left_object_order_not_guaranteed"
                ),
                "pptx_mode": options.pptx_mode.value,
            },
        )


__all__ = ["PptxParser"]
