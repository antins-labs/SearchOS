"""XLSX extraction that keeps formulas and cached values distinct."""

from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import Any

from searchos.local.models import (
    CellRangeLocator,
    DocumentKind,
    ParsedSegment,
    ParseResult,
    PipelineOptions,
)
from searchos.local.errors import LocalDocumentError, MissingDependencyError
from searchos.local.parsers.base import ParserCapability
from searchos.local.parsers.ooxml import inspect_ooxml_limits, validate_ooxml

_XLSX_CONTENT_TYPE = (
    b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
)


def _display(value: Any) -> str:
    if isinstance(value, (dt.date, dt.datetime, dt.time)):
        return value.isoformat()
    return str(value)


class XlsxParser:
    capability = ParserCapability(
        kind=DocumentKind.XLSX,
        extensions=(".xlsx",),
        locator_kinds=("cell_range",),
        dependencies=("openpyxl",),
    )

    def validate(self, path: Path) -> None:
        validate_ooxml(path, content_type_marker=_XLSX_CONTENT_TYPE)

    async def parse(self, path: Path, *, options: PipelineOptions) -> ParseResult:
        try:
            import openpyxl
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise MissingDependencyError(
                "XLSX 处理需要 local-docs 依赖；请安装 searchos[local-docs]"
            ) from exc

        inspect_ooxml_limits(
            path,
            max_uncompressed_bytes=options.max_archive_uncompressed_bytes,
            max_members=options.max_office_objects,
        )
        try:
            formulas = openpyxl.load_workbook(
                path,
                data_only=False,
                read_only=False,
                keep_links=False,
            )
            cached = openpyxl.load_workbook(
                path,
                data_only=True,
                read_only=False,
                keep_links=False,
            )
        except Exception as exc:  # noqa: BLE001
            raise LocalDocumentError(f"无法打开 XLSX：{exc}") from exc

        segments: list[ParsedSegment] = []
        sheet_metadata: list[dict[str, Any]] = []
        formulas_seen = 0
        external_formulas = 0
        total_cells = 0
        for sheet_number, worksheet in enumerate(formulas.worksheets, start=1):
            cached_sheet = cached[worksheet.title]
            dimensions = worksheet.max_row * worksheet.max_column
            total_cells += dimensions
            if total_cells > options.max_office_objects:
                raise LocalDocumentError(
                    f"XLSX 单元格扫描范围超过限制 {options.max_office_objects}"
                )
            hidden_rows = sum(1 for value in worksheet.row_dimensions.values() if value.hidden)
            hidden_columns = sum(
                1 for value in worksheet.column_dimensions.values() if value.hidden
            )
            merged_ranges = [str(value) for value in worksheet.merged_cells.ranges]
            for row_number in range(1, worksheet.max_row + 1):
                values: list[str] = []
                first_column: int | None = None
                last_column: int | None = None
                for column_number in range(1, worksheet.max_column + 1):
                    cell = worksheet.cell(row_number, column_number)
                    value = cell.value
                    if value is None:
                        continue
                    first_column = first_column or column_number
                    last_column = column_number
                    coordinate = cell.coordinate
                    if cell.data_type == "f":
                        formulas_seen += 1
                        formula = _display(value)
                        if "[" in formula and "]" in formula:
                            external_formulas += 1
                        cached_value = cached_sheet[coordinate].value
                        cached_text = (
                            _display(cached_value) if cached_value is not None else "<not cached>"
                        )
                        values.append(
                            f"{coordinate}: formula={formula}; cached={cached_text}"
                        )
                    else:
                        values.append(f"{coordinate}: {_display(value)}")
                if not values or first_column is None or last_column is None:
                    continue
                cell_range = (
                    f"{get_column_letter(first_column)}{row_number}:"
                    f"{get_column_letter(last_column)}{row_number}"
                )
                segments.append(
                    ParsedSegment(
                        id=f"X{sheet_number}-R{row_number}",
                        title=f"{worksheet.title} 第 {row_number} 行",
                        text="\n".join(values),
                        source_locators=[
                            CellRangeLocator(
                                sheet_id=str(worksheet.path),
                                sheet_name=worksheet.title,
                                range=cell_range,
                            )
                        ],
                    )
                )
            sheet_metadata.append(
                {
                    "sheet_id": str(worksheet.path),
                    "sheet_name": worksheet.title,
                    "state": worksheet.sheet_state,
                    "hidden_rows": hidden_rows,
                    "hidden_columns": hidden_columns,
                    "merged_ranges": merged_ranges,
                }
            )

        formulas.close()
        cached.close()
        warnings = []
        hidden_sheets = sum(item["state"] != "visible" for item in sheet_metadata)
        if hidden_sheets:
            warnings.append(f"保留了 {hidden_sheets} 个隐藏工作表的标记与内容")
        if external_formulas:
            warnings.append(f"检测到 {external_formulas} 个疑似外链公式；未访问外部链接")
        return ParseResult(
            kind=DocumentKind.XLSX,
            segments=segments,
            total_units=len(segments),
            covered_units=len(segments),
            warnings=warnings,
            parser_metadata={
                "sheets": sheet_metadata,
                "formulas": formulas_seen,
                "external_formulas": external_formulas,
                "recalculated": False,
                "external_links_loaded": False,
            },
        )


__all__ = ["XlsxParser"]
