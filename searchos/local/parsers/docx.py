"""DOCX paragraph and table extraction with OOXML-native locations."""

from __future__ import annotations

import zipfile
from pathlib import Path

from searchos.local.models import (
    DocumentKind,
    DocxBlockLocator,
    ParsedSegment,
    ParseResult,
    PipelineOptions,
)
from searchos.local.errors import MissingDependencyError
from searchos.local.parsers.base import ParserCapability
from searchos.local.parsers.ooxml import inspect_ooxml_limits, validate_ooxml

_DOCX_CONTENT_TYPE = (
    b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
)


def _table_text(table) -> tuple[str, list[DocxBlockLocator], int]:
    rows: list[str] = []
    locators: list[DocxBlockLocator] = []
    nested = 0
    seen_cells: set[int] = set()
    for row_number, row in enumerate(table.rows, start=1):
        values: list[str] = []
        for column_number, cell in enumerate(row.cells, start=1):
            cell_key = id(cell._tc)
            value = "\n".join(
                paragraph.text.strip() for paragraph in cell.paragraphs if paragraph.text.strip()
            )
            values.append(value.replace("|", "\\|").replace("\n", "<br>"))
            if cell_key not in seen_cells:
                seen_cells.add(cell_key)
                locators.append(
                    DocxBlockLocator(
                        part="word/document.xml",
                        path=("body",),
                        cell=f"row[{row_number}]/cell[{column_number}]",
                    )
                )
            nested += len(cell.tables)
        rows.append("| " + " | ".join(values) + " |")
    return "\n".join(rows), locators, nested


class DocxParser:
    capability = ParserCapability(
        kind=DocumentKind.DOCX,
        extensions=(".docx",),
        locator_kinds=("docx_block",),
        dependencies=("python-docx",),
    )

    def validate(self, path: Path) -> None:
        validate_ooxml(path, content_type_marker=_DOCX_CONTENT_TYPE)

    async def parse(self, path: Path, *, options: PipelineOptions) -> ParseResult:
        try:
            from docx import Document
            from docx.oxml.ns import qn
            from docx.table import Table
            from docx.text.paragraph import Paragraph
        except ImportError as exc:
            raise MissingDependencyError(
                "DOCX 处理需要 local-docs 依赖；请安装 searchos[local-docs]"
            ) from exc

        members = inspect_ooxml_limits(
            path,
            max_uncompressed_bytes=options.max_archive_uncompressed_bytes,
            max_members=options.max_office_objects,
        )
        document = Document(str(path))
        segments: list[ParsedSegment] = []
        nested_tables = 0
        paragraph_number = 0
        table_number = 0
        for child in document.element.body.iterchildren():
            if child.tag == qn("w:p"):
                paragraph_number += 1
                paragraph = Paragraph(child, document)
                content = paragraph.text.strip()
                if not content:
                    continue
                style_name = getattr(paragraph.style, "name", "") or ""
                is_heading = style_name.lower().startswith("heading")
                segments.append(
                    ParsedSegment(
                        id=f"D-P{paragraph_number}",
                        title=content[:200] if is_heading else f"段落 {paragraph_number}",
                        text=content,
                        source_locators=[
                            DocxBlockLocator(
                                part="word/document.xml",
                                path=("body", f"paragraph[{paragraph_number}]"),
                                paragraph_start=paragraph_number,
                                paragraph_end=paragraph_number,
                            )
                        ],
                        title_source="source" if is_heading else "generated",
                    )
                )
            elif child.tag == qn("w:tbl"):
                table_number += 1
                table = Table(child, document)
                content, locators, nested = _table_text(table)
                nested_tables += nested
                if content.strip():
                    locators = [
                        locator.model_copy(
                            update={
                                "path": ("body", f"table[{table_number}]"),
                                "cell": f"table[{table_number}]/{locator.cell}",
                            }
                        )
                        for locator in locators
                    ]
                    segments.append(
                        ParsedSegment(
                            id=f"D-T{table_number}",
                            title=f"表格 {table_number}",
                            text=content,
                            source_locators=locators,
                        )
                    )
            if len(segments) > options.max_office_objects:
                raise ValueError(f"DOCX 正文对象数超过限制 {options.max_office_objects}")

        warnings: list[str] = []
        with zipfile.ZipFile(path) as archive:
            document_xml = archive.read("word/document.xml")
        revisions = document_xml.count(b"<w:ins") + document_xml.count(b"<w:del")
        headers = sum(name.startswith("word/header") for name in members)
        footers = sum(name.startswith("word/footer") for name in members)
        comments = any(name == "word/comments.xml" for name in members)
        if revisions:
            warnings.append(f"检测到 {revisions} 个修订节点；当前只覆盖普通正文对象")
        if nested_tables:
            warnings.append(f"检测到 {nested_tables} 个嵌套表格；当前未展开其独立位置")
        if headers or footers:
            warnings.append(f"页眉 {headers} 个、页脚 {footers} 个当前未纳入正文表示")
        if comments:
            warnings.append("批注当前未纳入正文表示")
        return ParseResult(
            kind=DocumentKind.DOCX,
            segments=segments,
            total_units=len(segments),
            covered_units=len(segments),
            warnings=warnings,
            parser_metadata={
                "main_part": "word/document.xml",
                "paragraphs_seen": paragraph_number,
                "tables_seen": table_number,
                "revisions_uncovered": revisions,
                "nested_tables_uncovered": nested_tables,
                "headers_uncovered": headers,
                "footers_uncovered": footers,
                "comments_uncovered": comments,
            },
        )


__all__ = ["DocxParser"]
